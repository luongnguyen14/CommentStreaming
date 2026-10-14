#!/usr/bin/env python3
"""Task 3 - Fill the Work Breakdown Structure (WBS) template.

Copies the instructor's WBS workbook to deliverables/ and fills it with the
real project plan:

* E8 = team name, E9 = project topic (labels live in B8/B9 of the template);
* row 13 = the real calendar dates of D1-D20 (the template ships with a
  2021 sample row);
* rows 14-45 = the 32 plan rows: section titles in column B, work packages
  in column C, start/end dates, responsibility, deliverable and status;
  task texts are longer than the narrow hierarchy columns, so each row is
  merged (B:E for sections, C:E for work packages) - otherwise the vertical
  grid lines cut through the overflowing text in any PDF rendering;
* the Gantt cells K-AD are shaded for the day range each task spans, using
  the bar fill of the template's own sample row.

Nothing else is touched: fonts, borders, merged cells, the =TODAY() formula
in E7 and the Done/In-progress/Delayed dropdown on J14:J45 stay as-is.

Day 1 = 15/09/2026 (confirmed by the team on 14/10/2026), so D1-D20 maps to
15/09/2026 - 04/10/2026 and column K (D1) .. AD (D20).

Run:  python3 scripts/fill_wbs.py
"""

import copy
import datetime as dt
import os
import shutil
import sys

import openpyxl
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.environ.get("TEMPLATES_DIR", "/home/hadoop/capstone/templates")
TEMPLATE = os.path.join(
    TEMPLATES_DIR, "SIC_Big Data_Capstone Project_Work Breakdown Structure.xlsx"
)
OUT = os.path.join(PROJECT_ROOT, "deliverables", "WBS.xlsx")

TEAM_NAME = "Team 6"
PROJECT_TOPIC = (
    "SentimentGuard: Real-time Sentiment Classification of Live Chat Comments "
    "for Early Detection of Harmful Comment Spikes"
)

DAY1 = dt.date(2026, 9, 15)          # D1; D20 = 2026-10-04
GANTT_FIRST_COL = 11                 # column K = D1
FIRST_TASK_ROW = 14
LAST_TASK_ROW = 45

# (level, text, start day, end day, responsibility, deliverable, status)
# level "B" = section row, "C" = work package row.
# Status is the honest state on 14/10/2026 (Day 30 of the course calendar):
# only the planning work is finished, the rest is overdue or in progress.
TASKS = [
    ("B", "1. Project planning and management", None, None, "", "", "In-progress"),
    ("C", "Team setup and topic framing", 1, 2, "Ký", "docs/topic_decision.md", "Done"),
    ("C", "Action Plan preparation", 2, 3, "Ký, Dẳm", "deliverables/Action_Plan.docx", "Done"),
    ("C", "Data selection and source/licence verification", 2, 4, "Hạnh", "data/raw/SOURCE.md", "Done"),
    ("C", "Work Breakdown Structure (WBS) preparation", 4, 4, "Ký", "deliverables/WBS.xlsx", "In-progress"),
    ("C", "Preliminary presentation (Day 5)", 4, 5, "All members", "Proposal slides", "Delayed"),
    ("C", "WBS revision for review checkpoints", 9, 10, "Ký", "Updated WBS.xlsx (checks Day 9, 14, 19)", "Delayed"),

    ("B", "2. Data preparation", None, None, "", "", "Delayed"),
    ("C", "Data quality assessment (duplicates, rating distribution)", 6, 7, "Hạnh", "Quality report table", "Delayed"),
    ("C", "Shared text-cleaning module with unit tests", 6, 7, "Dẳm", "src/common/text_cleaning.py", "Delayed"),
    ("C", "Rating-to-sentiment mapping, dedup, fixed split", 7, 8, "Hạnh", "data/processed/ + eval_labels.parquet", "Delayed"),

    ("B", "3. Data ingestion pipeline (Kafka)", None, None, "", "", "Delayed"),
    ("C", "Kafka setup, topics and configuration", 8, 9, "Bằng", "kafka/ scripts + README", "Delayed"),
    ("C", "Replay producer (rate, burst, late records)", 9, 10, "Bằng", "src/producer/replay_producer.py", "Delayed"),
    ("C", "Ingestion throughput measurement", 10, 10, "Bằng", "outputs/tables/ingestion_*.csv", "Delayed"),

    ("B", "4. Data transformation (Spark Structured Streaming)", None, None, "", "", "Delayed"),
    ("C", "Streaming job: explicit schema, cleaning, DLQ", 10, 11, "Dẳm", "src/streaming/stream_job.py", "Delayed"),
    ("C", "Event-time windows and watermark aggregation", 11, 12, "Dẳm", "Parquet aggregates", "Delayed"),
    ("C", "Fault-tolerance test (restart mid-stream)", 12, 13, "Dẳm", "Test results table", "Delayed"),

    ("B", "5. Model training and evaluation", None, None, "", "", "Delayed"),
    ("C", "Spark ML pipeline; TF-IDF fitted on train only", 11, 12, "Hạnh", "src/training/train_model.py", "Delayed"),
    ("C", "Training and tuning with class weights", 12, 13, "Hạnh", "outputs/models/", "Delayed"),
    ("C", "Final test evaluation: macro-F1, per-class metrics", 13, 14, "Hạnh", "outputs/tables/metrics.csv", "Delayed"),

    ("B", "6. Query, visualization and dashboard", None, None, "", "", "Delayed"),
    ("C", "Spark SQL queries and aggregation tables", 14, 15, "Nguyên", "src/eda/", "Delayed"),
    ("C", "Figures and insight notes", 15, 16, "Nguyên", "outputs/figures/, insights.md", "Delayed"),
    ("C", "Streamlit dashboard and alert threshold", 16, 17, "Nguyên", "src/app/app.py", "Delayed"),

    ("B", "7. Prototype demonstration and delivery", None, None, "", "", "Delayed"),
    ("C", "Prototype showing presentation (Day 11)", 11, 11, "All members", "Prototype slides", "Delayed"),
    ("C", "Usage scenario and operating guide", 17, 18, "Dẳm, Ký", "docs/scenario.md", "Delayed"),
    ("C", "Final report (Day 18)", 17, 18, "All members", "deliverables/Final_Report.docx", "Delayed"),
    ("C", "Final presentation and live demo (Day 20)", 19, 20, "All members", "Final slides", "Delayed"),
]


def day_to_date(day):
    return DAY1 + dt.timedelta(days=day - 1)


def main():
    if not os.path.exists(TEMPLATE):
        sys.exit(f"Template not found: {TEMPLATE}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT)

    wb = openpyxl.load_workbook(OUT)
    ws = wb["WBS"]

    # --- header cells -----------------------------------------------------
    for cell_ref, value in (("E8", TEAM_NAME), ("E9", PROJECT_TOPIC)):
        ws[cell_ref] = value
        ws[cell_ref]._style = copy.copy(ws["E7"]._style)

    # --- row 13: real dates of D1-D20 ------------------------------------
    for day in range(1, 21):
        ws.cell(13, GANTT_FIRST_COL + day - 1).value = day_to_date(day)

    # --- task rows --------------------------------------------------------
    if len(TASKS) > LAST_TASK_ROW - FIRST_TASK_ROW + 1:
        sys.exit("Too many task rows for the template range J14:J45")

    bar_fill = copy.copy(ws["K14"].fill)      # the template's own Gantt bar
    empty_fill = copy.copy(ws["K18"].fill)    # a row with no bar

    for row in range(FIRST_TASK_ROW, LAST_TASK_ROW + 1):
        # Reset the whole row first so no sample data survives.
        for col in range(2, 31):
            ws.cell(row, col).value = None
        for col in range(GANTT_FIRST_COL, 31):
            ws.cell(row, col).fill = copy.copy(empty_fill)
        ws.cell(row, 10).value = None  # Task Status

    for i, (level, text, start, end, who, deliverable, status) in enumerate(TASKS):
        row = FIRST_TASK_ROW + i
        ws.cell(row, 2 if level == "B" else 3).value = text
        ws.cell(row, 10).value = status
        if start is None:               # section row: no schedule
            continue
        ws.cell(row, 6).value = day_to_date(start)   # F Start date
        ws.cell(row, 7).value = day_to_date(end)     # G End date
        ws.cell(row, 8).value = who                  # H Responsibility
        ws.cell(row, 9).value = deliverable          # I Deliverable
        for day in range(start, end + 1):            # Gantt bar
            ws.cell(row, GANTT_FIRST_COL + day - 1).fill = copy.copy(bar_fill)

    # Merge the hierarchy columns so long task names are not sliced by the
    # inner grid lines. Done last: a merged cell only accepts a value in its
    # top-left cell, which is where the text already sits.
    for i, (level, _text, *_rest) in enumerate(TASKS):
        row = FIRST_TASK_ROW + i
        ws.merge_cells(start_row=row, start_column=2 if level == "B" else 3,
                       end_row=row, end_column=5)

    wb.save(OUT)
    print(f"written: {OUT}")

    # --- verification dump ------------------------------------------------
    check = openpyxl.load_workbook(OUT)
    cws = check["WBS"]
    print(f"\nE7={cws['E7'].value!r}  E8={cws['E8'].value!r}  E9={str(cws['E9'].value)[:60]!r}")
    print("row 13:", " ".join(
        str(cws.cell(13, c).value)[:10] for c in range(11, 31)))
    print(f"\n{'row':>3}  {'task':<52} {'start':>10} {'end':>10} {'who':<12} {'status':<12} bar")
    for row in range(FIRST_TASK_ROW, LAST_TASK_ROW + 1):
        text = cws.cell(row, 2).value or cws.cell(row, 3).value or ""
        f = cws.cell(row, 6).value
        g = cws.cell(row, 7).value
        h = cws.cell(row, 8).value or ""
        j = cws.cell(row, 10).value or ""
        bars = "".join(
            "█" if cws.cell(row, c).fill.patternType else "·" for c in range(11, 31))
        print(f"{row:>3}  {text[:52]:<52} {str(f)[:10]:>10} {str(g)[:10]:>10} {h[:12]:<12} {j:<12} {bars}")

    print("\nmerges ok:", "B11:E13" in {str(m) for m in cws.merged_cells.ranges})
    dvs = cws.data_validations.dataValidation
    print("validation:", [(d.type, d.formula1, str(d.sqref)) for d in dvs])


if __name__ == "__main__":
    main()
