#!/usr/bin/env python3
"""Task 4 - Build the Day-5 proposal deck from the instructor's slide template.

Copies the Presentation Slide Template from TEMPLATES_DIR to
deliverables/Proposal_Slides.pptx and fills it in, keeping the template's
master, layouts, colours and fonts (SamsungOne 400 / Samsung Sharp Sans).

What the template provides and what we do with it:

* Slide 1 (Front Cover)  -> kept as the cover, title and team filled in.
* Slide 2 (3_Table of Contents2) -> dropped: a 5-slide proposal has no agenda slide.
* Slide 3 (2_Table of Contents layout) -> the content-slide design, duplicated
  for our four content slides. The yellow "speech bubble" on it is an
  instruction to us ("Please edit project name and page number in slide
  master") so it is removed from our copy.
* Slide 4 ('last' layout) -> dropped.

The template note is handled at the source: the footer text "Project Name"
lives in the LAYOUTS (bottom bar next to the page-number field), so we edit it
there - changing it once updates every slide. The page number itself is a
slide-number field and renders automatically (2..5).

Slide 3 of our deck ("Data and Architecture") carries the pipeline diagram,
drawn with PowerPoint shapes (theme navy, white fill) so it stays editable.

Speaker notes (~3 minutes per slide, 15 minutes total) are written for every
slide.

Run:  python3 scripts/fill_proposal_slides.py
"""

import copy
import os
import re
import shutil
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.environ.get("TEMPLATES_DIR", "/home/hadoop/capstone/templates")
TEMPLATE = os.path.join(
    TEMPLATES_DIR, "SIC_Big Data_Capstone Project_Presentation Slide Template.pptx"
)
OUT = os.path.join(PROJECT_ROOT, "deliverables", "Proposal_Slides.pptx")

TEAM_NAME = "Team 6"
MEMBERS_LINE_1 = "Vàng A Ký (Team Leader) · Nguyễn Đình Bằng · Vàng Thị Dẳm"
MEMBERS_LINE_2 = "Dương Thị Hạnh · Ngô Lương Nguyên"

SHORT_NAME = "SentimentGuard"
LONG_NAME = "Real-time Sentiment Classification of Live Chat Comments"

NAVY = RGBColor(0x14, 0x28, 0xA0)      # khớp dải màu xanh trên template
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# (number, subtitle, title, bullets, bullet box height (inches), notes)
SLIDES = [
    (
        "1.1",
        "Problem",
        "Problem and Technology for Good",
        [
            "Live chat never stops — thousands of comments arrive far faster than one person can read.",
            "Harmful spikes (harassment, hate, spam) are noticed only after the damage is done; "
            "manual moderation alone cannot keep up.",
            "Goal: detect a negative spike within seconds and warn the streamer while the session is still live.",
            "Who benefits: solo and small streamers, including educational and public-service channels.",
            "Technology for Good: protects community safety and creator mental health — self-hosted, "
            "no chat data leaves the machine.",
        ],
        2.4,
        "Good morning. We are Team 6. Our project is SentimentGuard — real-time sentiment classification "
        "of live chat comments, built to detect harmful comment spikes early.\n\n"
        "Let me start with the problem. A livestream chat is not a slow mailbox — it is a firehose. "
        "During a busy session, thousands of comments arrive, far faster than one person can read. "
        "When a harmful spike starts — harassment, hate, or coordinated spam — the streamer usually "
        "notices only after the damage is done. Commercial moderation tools exist, but they cost money, "
        "and the people who need them most — solo streamers and small educational or public-service "
        "channels — cannot afford a moderation team.\n\n"
        "So our goal is concrete: detect a spike of negative comments within seconds and warn the "
        "streamer while the session is still live, so they can slow the chat, enable moderation mode, "
        "or step away.\n\n"
        "Why this belongs to Technology for Good. First, the beneficiary is the public, not a company: "
        "the people who gain are small creators and their communities, and no product is being sold. "
        "Second, the cost of failure is personal: sustained harassment pushes creators to end streams, "
        "close their comments, or stop streaming altogether — that is a mental-health issue, not just "
        "a moderation backlog. Third, existing tooling leaves exactly this gap: large channels hire "
        "moderation teams or buy subscription services, while a solo streamer has neither, and "
        "self-hosted tools that watch comment sentiment over time and raise a spike warning are rare. "
        "And because our tool runs on the streamer's own machine, no chat data leaves it.\n\n"
        "To be concrete about what we are not claiming: we are not building a system that perfectly "
        "reads and judges every comment — that is beyond today's models, and beyond ours. We are "
        "building a system that notices when the conversation turns sharply negative, and tells the "
        "streamer in time to act.\n\n"
        "Next, I will show the data and the architecture we use to make this work.",
    ),
    (
        "1.2",
        "Data and Architecture",
        "Data and Architecture",
        [
            "Data: public IMDb reviews (CC0, about 1.0 million); 3 classes from 1–10 ratings.",
            "The stream is simulated: labelled live chat is unavailable, so we replay the data.",
            "Labels never enter the stream — stored separately, joined back by comment ID.",
            "Event time + watermark handle bursts; bad records go to a dead-letter queue.",
        ],
        1.6,
        "Now the data and the architecture. First, an honest statement about the data.\n\n"
        "Live chat with sentiment labels is not publicly available, and platform terms restrict keeping "
        "user comments. So we train the model on a public dataset — the IMDb review dataset under CC0, "
        "about one million reviews — and we derive three classes, negative, neutral and positive, from "
        "the 1-to-10 star ratings. We always call these rating-derived labels, never original labels, "
        "because a 3-star review is not literally a 'neutral' sentiment.\n\n"
        "The live stream itself is simulated: we replay the dataset through Kafka. This is a replay, not "
        "a real livestream, and we say so everywhere. In the stream, records carry bursts and deliberately "
        "late arrivals so that our event-time windows and watermarks are actually exercised, and bad "
        "records are not dropped silently — they go to a dead-letter queue with a reason.\n\n"
        "One rule we keep strictly: labels never travel with the stream. They are stored separately and "
        "joined back by comment ID only for evaluation. Training and streaming also share exactly one "
        "text-cleaning module, so they cannot drift apart.\n\n"
        "Architecture, left to right: a replay producer reads the dataset at a configurable rate and "
        "publishes JSON to the Kafka topic comments_raw. Spark Structured Streaming parses with an "
        "explicit schema, cleans the text, and aggregates per one-minute event-time windows with a "
        "watermark; invalid records branch to the dead-letter queue. A Spark ML pipeline — tokenizer, "
        "stop-word removal, n-grams and TF-IDF with a linear model — scores every comment. Results feed "
        "a small Streamlit dashboard that raises an alert when the negative ratio crosses a threshold.\n\n"
        "Why is there a Kafka broker in the middle, instead of just reading a file? Because we need a "
        "real streaming source: the producer can replay at a controlled rate, the buffer absorbs "
        "bursts, and the streaming job can be restarted without losing its place. The same topic could "
        "later carry real comments without changing anything else in the pipeline. Privacy is handled "
        "in the same path: user identifiers are replaced by salted hashes, personal fields are dropped, "
        "and the dashboard only ever shows aggregates, never individual comments.\n\n"
        "On measurement: the streaming job records how long each comment takes from arriving in the "
        "topic to being scored, so we can report end-to-end latency as p50, p95 and p99 — honest "
        "numbers instead of the phrase 'real time'.\n\n"
        "Everything runs locally on commodity hardware: Kafka, Spark and the dashboard.",
    ),
    (
        "1.3",
        "Outcome and Novelty",
        "Expected Outcome and Novelty",
        [
            "Near-real-time classification into negative / neutral / positive, with end-to-end latency "
            "reported (p50 / p95 / p99).",
            "Model quality reported as macro-F1 with per-class precision / recall / F1 — never accuracy alone.",
            "Novelty: detecting comment spikes, not just labelling comments — and one shared cleaning path "
            "for training and streaming, so no train/serve drift.",
            "Honest limits: rating-derived labels, domain shift (reviews are not chat), sarcasm; the stream "
            "is a replay — not a real livestream.",
            "Goal: a working prototype at TRL 5, validated in a simulated environment, released with "
            "reproducible scripts.",
        ],
        2.6,
        "What do we expect by the end, and what is new here?\n\n"
        "The deliverable is a prototype that classifies comments into negative, neutral and positive in "
        "near real time, and — this is the point — raises an alert within seconds when a harmful spike "
        "starts. We will report end-to-end latency at p50, p95 and p99, and model quality as macro-F1 "
        "with per-class precision, recall and F1. We deliberately avoid reporting accuracy alone, because "
        "with imbalanced data accuracy hides weak classes.\n\n"
        "Two things we believe are genuinely different from a plain sentiment demo. First, the product "
        "unit is a spike, not a label: we aggregate per time window and detect a rising negative ratio, "
        "which is what a streamer can actually act on. Second, training and streaming share exactly one "
        "cleaning implementation and one pipeline definition, so the classic train/serve drift cannot "
        "happen silently.\n\n"
        "Now the honest part, because we prefer clear limits to inflated claims. Our labels are derived "
        "from ratings, not from chat. Reviews are longer and better written than live chat, so there is "
        "a domain shift — a model that scores 0.85 on reviews will not score 0.85 on chat. Sarcasm "
        "remains hard. And our stream is a replay, not a real livestream; results from replayed data "
        "will always be labelled as such. Whatever we have not measured, we will mark as not yet verified "
        "rather than guess.\n\n"
        "How will we know it works? Three measurable checks, all reported as numbers rather than "
        "adjectives: the alert fires within a few seconds of a simulated spike; the model's weak "
        "classes are visible in per-class metrics instead of being hidden by accuracy; and throughput "
        "and latency are measured under bursts and across a restart. If a number is bad, it goes into "
        "the report as it is — we would rather show an honest weak result than an invented strong one.\n\n"
        "On maturity: a prototype assembled from real components and validated in a simulated "
        "environment is exactly TRL 5. Going further would need real users and real livestream data, "
        "which we deliberately do not touch for privacy and platform-terms reasons — and we will say "
        "this clearly in the report.",
    ),
    (
        "1.4",
        "Team and Plan",
        "Roles and Schedule",
        [
            "Ký — team lead, integration and testing; Bằng — Kafka ingestion and replay producer; "
            "Dẳm — streaming job, watermark, fault tolerance.",
            "Hạnh — data preparation and the Spark ML model; Nguyên — SQL queries, figures and the "
            "Streamlit dashboard with alerts.",
            "Phase 1 (Day 1–5) — planning: action plan, WBS, topic framing, and this proposal.",
            "Phase 2 (Day 6–11) — prototype: Kafka plus streaming, model trained, dashboard, prototype showing.",
            "Phase 3 (Day 12–20) — refinement: benchmarks, fault-tolerance tests, tuning, final report and "
            "presentation. WBS reviewed on Day 9, 14 and 19.",
        ],
        2.6,
        "Finally, who does what and when.\n\n"
        "Every member owns one span of the pipeline end to end, so responsibility is clear. Vàng A Ký "
        "leads the team and owns integration and testing. Nguyễn Đình Bằng owns Kafka: the broker, the "
        "topics, the replay producer and the ingestion throughput measurements. Vàng Thị Dẳm owns the "
        "streaming side: the Structured Streaming job, event-time windows, watermarks and the "
        "fault-tolerance test. Dương Thị Hạnh owns the data and the model: preparation, deduplication, "
        "the Spark ML pipeline and the final evaluation. Ngô Lương Nguyên owns everything the audience "
        "sees: SQL queries, figures and the Streamlit dashboard with the alert threshold.\n\n"
        "The schedule follows the three phases in our WBS. Phase 1, days one to five, is planning — and "
        "this presentation is its last item. Phase 2, days six to eleven, builds the prototype: Kafka "
        "and streaming, training data, the model, and the dashboard, finishing with the prototype "
        "showing. Phase 3, days twelve to twenty, is refinement: throughput benchmarks, the restart "
        "test, model tuning, the usage scenario, and finally the report and presentation.\n\n"
        "The WBS is reviewed three times, on Day 9, 14 and 19, and we update task status as we go — "
        "the workbook you have reflects the real state, not an ideal one.\n\n"
        "Two things about how we work as a team, because the evaluation also looks at project "
        "management. First, every deliverable in this plan is committed to git with a dated progress "
        "log, so the history shows what was done and when — the WBS status you see is the real status, "
        "including the parts that slipped. Second, we have a documented answer to the resource risk: "
        "the machine has 7 GB of RAM, so we cap Spark's memory, limit partitions and the replay rate, "
        "and we measure how far the machine can actually go instead of assuming.\n\n"
        "To sum up: a clear problem, an honest data strategy, a pipeline we can build in the time we "
        "have, and results that will be reported exactly as measured.\n\n"
        "That is our proposal. Thank you — we are happy to take questions.",
    ),
]


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def set_para_text(para, text, size=None, color=None, bold=None):
    """Replace a paragraph's text, keeping the first run's inherited formatting."""
    runs = para.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        runs = [para.add_run()]
        runs[0].text = text
    for r in runs:
        if size is not None:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if bold is not None:
            r.font.bold = bold


def set_para_text_safe_digits(para, text):
    """Like set_para_text, but writes each run of digits as its own run with an
    explicit latin typeface.

    Why: on the cover layout LibreOffice shapes 'Samsung Sharp Sans' (the
    template's obfuscated brand font) wrongly for a lowercase letter followed
    by a space and a digit — 'Team 6' comes out as 'Tea 6' plus a .notdef box,
    which also lands in the exported PDF. A run boundary stops that
    substitution, so digits get their own run."""
    pieces = [s for s in re.split(r"(\d+)", text) if s]
    if not para.runs:
        para.add_run()
    anchor = para.runs[0]._r
    donor = copy.deepcopy(anchor)
    for extra in list(para.runs[1:]):
        extra._r.getparent().remove(extra._r)

    def write(r_el, piece):
        r_el.find(qn("a:t")).text = piece
        if not piece.isdigit():
            return
        rPr = r_el.find(qn("a:rPr"))
        if rPr is None:
            rPr = etree.Element(qn("a:rPr"))
            r_el.insert(0, rPr)
        latin = rPr.find(qn("a:latin"))
        if latin is None:
            latin = etree.SubElement(rPr, qn("a:latin"))
        latin.set("typeface", "Arial")

    write(anchor, pieces[0])
    prev = anchor
    for piece in pieces[1:]:
        r_el = copy.deepcopy(donor)
        write(r_el, piece)
        prev.addnext(r_el)
        prev = r_el


def set_cover_title(shape, line1, line2, line2_size=22):
    """Cover title: one paragraph, two runs separated by a line break (as the template)."""
    tf = shape.text_frame
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    p = tf.paragraphs[0]
    for child in list(p._p):               # drop old runs and <a:br/> line breaks
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p._p.remove(child)
    r1 = p.add_run()
    r1.text = line1
    p.add_line_break()
    r2 = p.add_run()
    r2.text = line2
    r2.font.size = Pt(line2_size)


def set_shape_para(shape, texts_with_sizes, color=None, bold=None):
    """Overwrite a (single-paragraph) placeholder with one paragraph + text."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    set_para_text(first, texts_with_sizes[0], size=texts_with_sizes[1],
                  color=color, bold=bold)


def materialize_geometry(slide, shape, height_in=None):
    """Copy position/size from the layout placeholder into the slide shape.

    The slide placeholders inherit their geometry from the layout; writing only
    `shape.height` makes python-pptx create an <a:xfrm> with cx=0, and the
    shape then renders with zero width (invisible text). Materializing the
    layout's left/top/width first avoids that.
    """
    idx = shape.placeholder_format.idx
    for lay_sh in slide.slide_layout.placeholders:
        if lay_sh.placeholder_format.idx == idx:
            shape.left = lay_sh.left
            shape.top = lay_sh.top
            shape.width = lay_sh.width
            break
    if height_in is not None:
        shape.height = Inches(height_in)


def fill_bullets(shape, bullets, size=15):
    """Replace the bullet placeholder's paragraphs, cloning the first paragraph."""
    tf = shape.text_frame
    tf.word_wrap = True
    donor = copy.deepcopy(tf.paragraphs[0]._p)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for b in bullets:
        tf._txBody.append(copy.deepcopy(donor))
        set_para_text(tf.paragraphs[-1], b, size=size)


def find_placeholder(slide, idx):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == idx:
            return sh
    raise KeyError(f"placeholder idx={idx} not found")


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    el = list(sldIdLst)[index]
    rId = el.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(el)


def duplicate_slide(prs, index):
    """Append a copy of slide `index` (plain placeholder shapes only)."""
    src = prs.slides[index]
    dest = prs.slides.add_slide(src.slide_layout)
    for shp in list(dest.shapes):          # drop the layout-cloned placeholders
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:                 # clone the source slide's own shapes
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    return dest


def add_box(slide, x, y, w, h, title, detail=None, dashed=False):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.14
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = NAVY
    box.line.width = Pt(1.25)
    if dashed:
        box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    set_para_text(tf.paragraphs[0], title, size=11, color=NAVY, bold=True)
    if detail:
        p = tf.add_paragraph()
        set_para_text(p, detail, size=8.5, color=GREY)
    return box


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = NAVY
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "med", "len": "med"}))
    return conn


def add_caption(slide, x, y, w, text, size=9.5):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.3))
    tb.text_frame.word_wrap = True
    set_para_text(tb.text_frame.paragraphs[0], text, size=size, color=GREY)
    return tb


def add_pipeline_diagram(slide):
    """Pipeline diagram on the Data and Architecture slide (editable shapes)."""
    add_caption(slide, 0.57, 4.05, 9.6,
                "Simulated live stream — the dataset is replayed through Kafka (not a real livestream):")

    y, h = 4.40, 0.95
    boxes = [
        (0.57, 1.70, "Replay producer", "rate · bursts · late records"),
        (2.62, 1.70, "Kafka topic", "comments_raw (JSON)"),
        (4.67, 1.85, "Spark Structured Streaming", "schema · event-time windows · watermark"),
        (6.87, 1.70, "Spark ML pipeline", "TF-IDF + linear model → sentiment"),
        (8.92, 1.70, "Streamlit dashboard", "aggregates + spike alert"),
    ]
    for x, w, t, d in boxes:
        add_box(slide, x, y, w, h, t, d)
    for x in (2.27, 4.32, 6.52, 8.57):
        add_arrow(slide, x, y + h / 2, x + 0.35, y + h / 2)

    # dead-letter queue branch under the streaming box
    add_arrow(slide, 5.59, y + h, 5.59, y + h + 0.45)
    add_box(slide, 4.67, y + h + 0.45, 1.85, 0.62,
            "Dead-letter queue", "invalid records + reason", dashed=True)

    add_caption(slide, 0.57, 6.62, 9.6,
                "Labels stay outside the stream: stored separately and joined by comment ID only for "
                "evaluation.", size=9.5)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def main():
    if not os.path.exists(TEMPLATE):
        sys.exit(f"Template not found: {TEMPLATE}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT)

    prs = Presentation(OUT)

    # --- 1. Template note: project name lives in the LAYOUTS' footer ---------
    for layout in prs.slide_layouts:
        for sh in layout.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "Project Name":
                set_para_text(sh.text_frame.paragraphs[0], SHORT_NAME)
            # cover layout prompt text follows the deck
            if sh.is_placeholder and sh.placeholder_format.idx == 0 \
                    and "Project" in sh.text_frame.text:
                set_cover_title(sh, SHORT_NAME, LONG_NAME)
            if sh.is_placeholder and sh.placeholder_format.idx == 11 \
                    and sh.text_frame.text.strip() == "Team Name":
                set_para_text(sh.text_frame.paragraphs[0], TEAM_NAME)

    # --- 2. Cover ------------------------------------------------------------
    cover = prs.slides[0]
    set_cover_title(find_placeholder(cover, 0), SHORT_NAME, LONG_NAME)

    body = find_placeholder(cover, 11)
    tf = body.text_frame
    set_para_text_safe_digits(tf.paragraphs[0], TEAM_NAME)
    p = tf.add_paragraph()
    set_para_text(p, MEMBERS_LINE_1, size=13)
    p = tf.add_paragraph()
    set_para_text(p, MEMBERS_LINE_2, size=13)
    materialize_geometry(cover, body, height_in=1.1)

    # --- 3. Content slides ----------------------------------------------------
    content = prs.slides[2]
    for sh in list(content.shapes):        # remove the instruction speech bubble
        if "Please edit project name" in (sh.text_frame.text if sh.has_text_frame else ""):
            sh._element.getparent().remove(sh._element)

    # Duplicate the content slide FIRST: deleting a slide only drops its
    # relationship, the part stays in the package, and add_slide() would then
    # pick a part name that is still in use (two slides ending up in one
    # slideN.xml and showing identical content).
    for _ in range(len(SLIDES) - 1):       # 4 content slides total
        duplicate_slide(prs, 2)

    # drop the agenda slide and the closing slide
    delete_slide(prs, 3)                   # 'last'
    delete_slide(prs, 1)                   # 3_Table of Contents2

    for i, (num, subtitle, title_text, bullets, height, notes) in enumerate(SLIDES):
        slide = prs.slides[i + 1]
        set_shape_para(find_placeholder(slide, 0), (title_text, None))
        set_shape_para(find_placeholder(slide, 10), (num, None))
        set_shape_para(find_placeholder(slide, 11), (subtitle, None))
        set_shape_para(find_placeholder(slide, 12), ("01", None))
        bullet_shape = find_placeholder(slide, 15)
        materialize_geometry(slide, bullet_shape, height_in=height)
        fill_bullets(bullet_shape, bullets, size=15)
        slide.notes_slide.notes_text_frame.text = notes

    add_pipeline_diagram(prs.slides[2])

    # cover notes (slide order: cover, 4 content slides)
    prs.slides[0].notes_slide.notes_text_frame.text = (
        "Good morning. We are Team 6, and this is SentimentGuard — real-time sentiment classification "
        "of live chat comments, built to detect harmful comment spikes early.\n\n"
        "The team: Vàng A Ký, our team leader, who coordinates the project and integration testing. "
        "Nguyễn Đình Bằng, our Kafka and data engineer. Vàng Thị Dẳm, our streaming engineer. "
        "Dương Thị Hạnh, our machine learning engineer. And Ngô Lương Nguyên, our data analyst and "
        "dashboard developer.\n\n"
        "The name SentimentGuard says what the project does: sentiment analysis used as a guard for "
        "the community. The point is to protect people — the streamer and the audience — not to sell "
        "anything.\n\n"
        "In the next four slides we will cover the problem and why it matters for the community, the "
        "data and the architecture, what we expect to deliver and what is genuinely new, and finally "
        "the roles and the schedule.\n\n"
        "One note before we start, because honesty is a rule for us: our working stream is a replay of "
        "a public dataset through Kafka — it is not a real livestream — and wherever a result comes "
        "from this simulation we will label it as simulation."
    )

    prs.save(OUT)
    print(f"written: {OUT}")

    # --- verification dump ----------------------------------------------------
    check = Presentation(OUT)
    print(f"\nso slide: {len(check.slides)}")
    for i, slide in enumerate(check.slides):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.replace("\n", " / ")[:70])
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        print(f"\n--- slide {i + 1} (layout={slide.slide_layout.name}) ---")
        for t in texts:
            print(f"    {t!r}")
        print(f"    notes: {len(notes.split())} words")

    footers = [sh.text_frame.text for lay in check.slide_layouts for sh in lay.shapes
               if sh.has_text_frame and sh.text_frame.text.strip() == "Project Name"]
    print("\nfooter con sot 'Project Name':", len(footers))


if __name__ == "__main__":
    main()
